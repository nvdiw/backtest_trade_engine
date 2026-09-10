"""Shared readable headers and directional cues for generated Excel reports."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import CellIsRule


def style_report(workbook):
    for sheet in workbook.worksheets:
        max_row = sheet.max_row
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = '1:1'
        sheet.freeze_panes = ('C2' if sheet.title in ('Quick Compare', 'Stage Observations',
                                                     'Direction Leaders') else sheet.freeze_panes or 'A2')
        sheet.row_dimensions[1].height = 46
        for cell in sheet[1]:
            name = str(cell.value or '').lower()
            color = ('217346' if 'long' in name else 'B54758' if 'short' in name
                     else '75549B' if 'drawdown' in name or 'liquidation' in name
                     else '17365D')
            cell.fill = PatternFill('solid', fgColor=color)
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
            cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
            sheet.column_dimensions[cell.column_letter].width = max(
                16, sheet.column_dimensions[cell.column_letter].width or 0)
            if max_row < 2:
                continue
            if any(word in name for word in ('profit', 'net_return', 'expectancy')) and not any(
                    word in name for word in ('factor', 'months', 'ratio', 'target')):
                span = f'{cell.column_letter}2:{cell.column_letter}{max_row}'
                for operator, fill, font in [('lessThan', 'FCE4D6', '9C0006'),
                                              ('greaterThan', 'E2F0D9', '006100')]:
                    sheet.conditional_formatting.add(span, CellIsRule(
                        operator=operator, formula=['0'],
                        fill=PatternFill('solid', fgColor=fill), font=Font(color=font)))
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
