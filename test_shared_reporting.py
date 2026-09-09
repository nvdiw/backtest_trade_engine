import contextlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import ma_strategy
import pulse_strategy
import optimize
from market_data import clear_market_data_cache
from optimize_commands import expand_arguments
from campaign_reporting import publish_campaign
from test_pulse_strategy import candles
from trade_engine import TradeEngine


class SharedReportingTests(unittest.TestCase):
    def tearDown(self):
        clear_market_data_cache()
        TradeEngine.load_market_data.cache_clear()
        pulse_strategy._FEATURE_CACHE.clear()

    def test_ma_pulse_and_new_market_interval_have_same_artifacts(self):
        expected = {'result.json', 'params.json', 'run_summary.csv', 'run.log', 'chart.png',
                    'trades/data_orders.csv', 'trades/data_orders.xlsx', 'trades/data_orders_summary.csv',
                    'trades/data_orders_monthly_summary.csv', 'trades/data_orders_monthly_targets.csv',
                    'monthly/monthly_data_orders.csv'}
        sheet_sets, column_sets = [], []
        with tempfile.TemporaryDirectory() as temp, patch.dict(os.environ, {'BTE_DATA_FILE': '', 'BTE_TIMEFRAME': 'auto'}):
            for module, interval, symbol in ((pulse_strategy, '1m', 'BTC'), (ma_strategy, '15m', 'BTC'),
                                               (pulse_strategy, '5m', 'ETH')):
                label = f'{module.__name__}_{interval}'
                path = Path(temp) / f'{symbol}_{interval}_candles.csv'
                frame = candles(600)
                delta = pd.Timedelta(minutes=int(interval[:-1]))
                frame['Open time'] = pd.date_range('2025-01-01', periods=len(frame), freq=delta)
                frame['Close time'] = frame['Open time'] + delta - pd.Timedelta(milliseconds=1)
                frame.to_csv(path, index=False)
                out = Path(temp) / label
                function = module.pulse_strategy if module is pulse_strategy else module.ma_strategy
                with contextlib.redirect_stdout(io.StringIO()):
                    result = function({'optimize': False}, start=0, end=len(frame),
                                      data_file=path, timeframe='auto', output_dir=out, verbose=False,
                                      show_chart=False, chart_file=out / 'chart.png')
                self.assertTrue(expected <= {str(p.relative_to(out)).replace('\\', '/') for p in out.rglob('*') if p.is_file()})
                self.assertEqual(result['timeframe'], interval)
                self.assertEqual(result['symbol'], symbol)
                self.assertEqual(result['report_schema_version'], 1)
                self.assertGreater((out / 'chart.png').stat().st_size, 1000)
                book = load_workbook(out / 'trades/data_orders.xlsx')
                sheet_sets.append(book.sheetnames)
                book.close()
                column_sets.append(pd.read_csv(out / 'trades/data_orders.csv').columns.tolist())
                persisted = json.loads((out / 'result.json').read_text())
                self.assertEqual(persisted['timeframe'], interval)
                self.assertIn('chart_file', persisted)
            self.assertTrue(all(sheets == sheet_sets[0] for sheets in sheet_sets))
            self.assertTrue(all(columns == column_sets[0] for columns in column_sets))

    def test_campaign_without_winner_still_has_csv_excel_log(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            state = dict(status='completed', cycle=201, cycles_completed=200, total_evaluations=2,
                         config={'strategy': 'pulse_strategy:pulse_strategy', 'timeframe': '1m',
                                 'minimum_trades': 30, 'maximum_allowed_drawdown': 25})
            (path / 'auto_state.json').write_text(json.dumps(state))
            stage = path / 'cycles/cycle_000001'
            stage.mkdir(parents=True)
            pd.DataFrame([dict(candidate_id='a', closed_trades=5, total_profit_percent=2,
                               maximum_drawdown=-1, objective_score=None, error=''),
                          dict(candidate_id='b', closed_trades=0, total_profit_percent=0,
                               maximum_drawdown=0, objective_score=None, error='')]).to_csv(stage / 'validation_results.csv', index=False)
            result = publish_campaign(path, rebuild=True)
            self.assertEqual(result['outcome'], 'NO_QUALIFIED_FINALIST')
            self.assertIsNone(result['recommended_params'])
            stages = pd.read_csv(path / 'stage_summary.csv')
            self.assertEqual(stages.iloc[0]['no_trades'], 1)
            self.assertEqual(stages.iloc[0]['insufficient_trades'], 1)
            self.assertTrue((path / 'campaign_report.xlsx').is_file())
            self.assertTrue((path / 'campaign.log').is_file())

    def test_short_commands_and_duration_trade_gate(self):
        self.assertEqual(expand_arguments(['pulse', '--cycles', '200']),
                         ['--strategy', 'pulse', '--auto', '--cycles', '200'])
        self.assertEqual(expand_arguments(['resume', 'folder']), ['--resume', 'folder'])
        self.assertEqual(optimize.duration_trade_requirement(30, 90, 270), 10)
        self.assertEqual(optimize.duration_trade_requirement(30, 30, 270), 4)
        self.assertEqual(optimize.duration_trade_requirement(0, 30, 270), 0)


if __name__ == '__main__':
    unittest.main()
