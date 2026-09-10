import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from monthly_reporting import monthly_report
from trade_engine import TradeEngine, AccountState


class MonthlyReportingTests(unittest.TestCase):
    def test_ten_months_eight_targets_met(self):
        summary, rows = monthly_report('2025-01-01', '2025-11-01', [.08] * 8 + [.03, -.01], 8)
        self.assertEqual(summary['calendar_months'], 10)
        self.assertEqual(summary['monthly_target_met_months'], 8)
        self.assertEqual(summary['monthly_target_missed_months'], 2)
        self.assertEqual(summary['monthly_target_success_percent'], 80)
        self.assertEqual(summary['monthly_positive_months'], 9)
        self.assertEqual(summary['backtest_days'], 304)
        self.assertEqual(summary['monthly_full_months'], 10)
        self.assertEqual(rows[-1]['target_status'], 'MISSED')

    def test_partial_and_missing_months_are_explicit(self):
        summary, rows = monthly_report('2025-01-15', '2025-04-10', [.1, 0], 9,
                                      ['2025-01', '2025-03'])
        self.assertEqual(summary['monthly_partial_months'], 2)
        self.assertEqual(summary['monthly_target_unknown_months'], 2)
        self.assertEqual(summary['monthly_target_missed_months'], 1)
        self.assertEqual(rows[1]['target_status'], 'UNKNOWN')
        self.assertTrue(rows[0]['partial_month'])

    def test_csv_and_workbook_include_monthly_result_with_no_trades(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = TradeEngine(optimize=False, verbose=False, output_dir=directory)
            result = engine.finalize_account(AccountState(balance=1000), first_balance=1000,
                ending_mark_price=100, start_time='2025-01-01', end_time='2025-03-01',
                monthly_returns=[0., .08], monthly_profit_target_percent=8)
            self.assertEqual(result['monthly_target_met_months'], 1)
            path = Path(directory) / 'trades/data_orders.csv'
            summary = pd.read_csv(path).iloc[-1]
            self.assertEqual(summary['monthly_target_met_months'], 1)
            self.assertEqual(summary['monthly_target_missed_months'], 1)
            self.assertEqual(summary['backtest_days'], 59)
            detail = pd.read_csv(path.with_name('data_orders_monthly_targets.csv'))
            self.assertEqual(detail['target_status'].tolist(), ['MISSED', 'MET'])
            book = load_workbook(path.with_suffix('.xlsx'))
            self.assertEqual(book.sheetnames[:3], ['Dashboard', 'Monthly Summary', 'Monthly Goals'])
            self.assertEqual(book['Monthly Goals'].max_row, 3)
            book.close()

    def test_log_counts_calendar_losses_and_eight_percent_without_loss_stop(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = TradeEngine(optimize=False, verbose=False, output_dir=directory)
            result = engine.finalize_account(AccountState(balance=1000), first_balance=1000,
                ending_mark_price=100, start_time='2025-01-01', end_time='2025-11-01',
                monthly_returns=[.08] * 8 + [.03, -.01], monthly_profit_target_percent=9,
                report_metadata={'monthly_loss_close_filter': False})
            log = (Path(directory) / 'run.log').read_text(encoding='utf-8')
            self.assertIn('Total calendar months: 10', log)
            self.assertIn('Months with profit >= 8%: 8', log)
            self.assertIn('Losing months (net return < 0%): 1', log)
            self.assertIn('Months below 8%: 2', log)
            self.assertIn('Monthly loss stop enabled: false', log)
            self.assertEqual(result['monthly_target_met_months'], 0)  # Separate 9% report target.


if __name__ == '__main__':
    unittest.main()
