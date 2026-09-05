"""Shared readable chart layout for optimizer, audit and trade workbooks."""

import math


def add_report_chart(workbook, chart):
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties

    sheet = workbook['Charts'] if 'Charts' in workbook.sheetnames else workbook.create_sheet('Charts')
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A2'
    chart.width, chart.height = 24, 11
    chart.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1000)))])
    points = 0
    for series in chart.series:
        reference = getattr(getattr(series, 'val', None), 'numRef', None)
        if reference and reference.f:
            from openpyxl.utils.cell import range_boundaries
            _, first, _, last = range_boundaries(reference.f.split('!')[-1])
            points = max(points, last-first+1)
    # Exact values remain in the data tables. Dense multi-series labels are
    # deliberately suppressed rather than drawn on top of adjacent bars.
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = len(chart.series) == 1 and points <= 6
    chart.dLbls.showSerName = False
    chart.dLbls.showCatName = False
    chart.dLbls.showLegendKey = False
    chart.dLbls.numFmt = getattr(chart.y_axis, 'numFmt', None) or '#,##0.0'
    if hasattr(chart.x_axis, 'tickLblSkip'):
        chart.x_axis.tickLblSkip = max(1, math.ceil(points / 12))
        chart.x_axis.tickMarkSkip = chart.x_axis.tickLblSkip
        chart.x_axis.tickLblPos = 'low'
    if chart.legend:
        chart.legend.position = 'b'
        chart.legend.overlay = False
    chart_number = len(sheet._charts)
    anchor_row = 2 + chart_number * 24
    for row in range(anchor_row, anchor_row+24):
        sheet.row_dimensions[row].height = 16
    sheet.add_chart(chart, f'A{anchor_row}')
