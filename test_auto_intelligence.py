import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

import optimize
from excel_charts import add_report_chart


class AutoIntelligenceTests(unittest.TestCase):
    def test_directional_schedule_and_seed_inheritance(self):
        args = optimize.build_parser().parse_args(['--auto', '--staged', '--directional'])
        phases = optimize._strategy_staged_phases(args)
        self.assertEqual(len(phases), 11)
        self.assertEqual(phases[:2], (('long_signal', 'long_signal'), ('short_signal', 'short_signal')))
        self.assertEqual(phases[-1], ('portfolio', 'portfolio'))
        seeds = optimize._staged_seed_data([{'effective_params': {'ma_50_period': 60,
                                                                 'short_ma_50_period': None},
                                            'robust_score': 10}],
                                          {'short_ma_50_period': [40, 50, 60]}, {})
        self.assertEqual(seeds[0]['params']['short_ma_50_period'], 60)
        self.assertIn('rsi_short_tp_pct', optimize.PARAMETER_PROFILES['short_rsi'])
        self.assertNotIn('rsi_long_tp_pct', optimize.PARAMETER_PROFILES['short_rsi'])

    def test_weak_surrogate_reserves_more_random_exploration(self):
        history = [{'params': {'x': i}, 'learning_score': 1.0} for i in range(40)]
        candidates = [{'x': i} for i in range(40, 80)]
        features = {'surrogate_min_samples': 4, 'surrogate_trees': 4, 'surrogate_max_samples': 100}
        selected, metadata = optimize._select_surrogate_candidates(
            candidates, history, 12, ('x',), {'x': list(range(80))}, features, 42)
        self.assertEqual(len(selected), 12)
        self.assertEqual(metadata['validation_rank_correlation'], 0)
        self.assertAlmostEqual(metadata['selection_mix']['random'], .35)
        self.assertAlmostEqual(sum(metadata['selection_mix'].values()), 1)

    def test_snapshot_workbook_reuses_evaluated_windows(self):
        args = optimize.build_parser().parse_args(['--random-audit-tests', '2', '--random-audit-top', '1', '-w', '1'])
        windows = [dict(window_id=i, tier='historical', months=6, start=0, end=100,
                        start_label='2024-01-01', end_label='2024-07-01') for i in (1,2)]
        result = dict(score=10, total_profit_percent=20, closed_trades=50,
                      maximum_drawdown=-5, liquidations=0, monthly_returns=[.1], trade_profits=[10])
        def evaluate(task):
            index, candidate, window, params, start, end, _ = task
            return index, candidate, window, params, start, end, result, .01, None
        with tempfile.TemporaryDirectory() as directory:
            with patch.object(optimize, '_generate_random_audit_windows', return_value=windows), \
                 patch.object(optimize, '_evaluate_random_window_task', side_effect=evaluate) as worker:
                optimize._run_random_window_audit(args, [{'candidate_id': 'winner', 'effective_params': {'balance':1000}}],
                                                  directory, 1, 100)
                self.assertEqual(worker.call_count, 2)
            book = load_workbook(Path(directory) / 'random_window_report.xlsx')
            self.assertEqual(book['Windows'].max_row, 3)
            self.assertIn('Charts', book.sheetnames)
            book.close()

    def test_dense_charts_have_separate_nonoverlapping_anchors(self):
        book = Workbook()
        data = book.active
        data.append(['candidate', 'long', 'short'])
        for i in range(30):
            data.append([f'candidate-{i:06d}', i, i+1])
        for _ in range(2):
            chart = BarChart()
            chart.add_data(Reference(data, min_col=2, max_col=3, min_row=1, max_row=31), titles_from_data=True)
            chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=31))
            add_report_chart(book, chart)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'charts.xlsx'
            book.save(path)
            loaded = load_workbook(path)
            charts = loaded['Charts']._charts
            self.assertFalse(charts[0].dLbls.showVal)
            self.assertEqual(charts[0].x_axis.tickLblSkip, 3)
            self.assertGreater(charts[1].anchor._from.row, charts[0].anchor._from.row)
            self.assertEqual(charts[0].legend.position, 'b')
            self.assertFalse(charts[0].legend.overlay)
            loaded.close()


if __name__ == '__main__':
    unittest.main()
