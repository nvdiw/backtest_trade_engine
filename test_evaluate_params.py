import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from evaluate_params import generate_windows, summarize, write_workbook


class FakeSource:
    def __init__(self):
        self.times = pd.Series(pd.date_range('2020-01-01', '2023-01-01', freq='D', inclusive='left'))

    def coverage(self):
        return {'first_candle': '2020-01-01', 'end_exclusive': '2023-01-01'}

    def open_times(self):
        return self.times

    def resolve_index(self, value):
        return int(self.times.searchsorted(pd.Timestamp(value)))

    def format_bound(self, index):
        return str(self.times.iloc[index] if index < len(self.times) else pd.Timestamp('2023-01-01'))


class FixedEvaluationTests(unittest.TestCase):
    def test_unique_reproducible_calendar_windows(self):
        source = FakeSource()
        windows = generate_windows(source, None, 'latest', 1000, 6, 12, 42)
        self.assertEqual(windows, generate_windows(source, None, 'latest', 1000, 6, 12, 42))
        self.assertEqual(len({(w['start'], w['end']) for w in windows}), 1000)
        for w in windows:
            self.assertLessEqual(w['end'], len(source.times))
            self.assertEqual(pd.Timestamp(w['end_exclusive']),
                             pd.Timestamp(w['start_date']) + pd.DateOffset(months=w['months']))

    def test_short_and_outside_ranges_rejected(self):
        for start in ('2019-01-01', '2022-12-01'):
            with self.assertRaises(ValueError):
                generate_windows(FakeSource(), start, 'latest', 10, 6, 12, 42)

    def record(self, **metrics):
        return dict(window_id=1, start_date='2020-01-01', end_exclusive='2021-01-01', months=12,
                    error=None, result=dict(total_profit_percent=20, maximum_drawdown=-10,
                                            closed_trades=30, liquidations=0, **metrics))

    def test_errors_and_missing_metrics_never_pass(self):
        records = [self.record(), dict(window_id=2, error='failed', result={}),
                   dict(window_id=3, result={'total_profit_percent': 100})]
        summary, frame = summarize(records, 3, 5, 30, .3, .3)
        self.assertEqual(summary['status'], 'FAIL')
        self.assertEqual(summary['failed_windows'], 2)
        self.assertAlmostEqual(summary['positive_window_ratio'], 1/3)
        self.assertEqual(frame.window_pass.tolist(), [True, False, False])

    def test_risk_gate_and_incomplete_execution(self):
        record = self.record()
        record['result']['liquidations'] = 1
        summary, _ = summarize([record], 1, 5, 30, .7, .7)
        self.assertEqual(summary['status'], 'FAIL')
        summary, _ = summarize([self.record()], 2, 5, 30, .3, .3)
        self.assertFalse(summary['gates']['complete'])

    def test_workbook_has_full_window_and_series_data(self):
        records = [self.record(monthly_returns=[.1, -.02], trade_profits=[10, -2])]
        summary, frame = summarize(records, 1, 5, 30, .7, .7)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.xlsx'
            write_workbook(path, frame, summary, {'balance': 1000}, {'seed': 42}, records)
            book = load_workbook(path)
            self.assertEqual(book['Windows'].max_row, 2)
            self.assertEqual(book['Monthly Returns'].max_row, 3)
            self.assertEqual(book['Trade Profits 1'].max_row, 3)
            self.assertEqual(book['Monthly Returns']['C2'].value, .1)
            self.assertEqual(book['Windows'].freeze_panes, 'A2')
            self.assertEqual(len(book['Distribution']._charts), 1)
            book.close()


if __name__ == '__main__':
    unittest.main()
