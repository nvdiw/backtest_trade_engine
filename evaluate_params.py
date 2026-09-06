"""Evaluate ONE frozen configuration on reproducible random calendar windows."""

from excel_charts import add_report_chart

import argparse
import json
import math
import multiprocessing
import random
import time
from pathlib import Path

import pandas as pd

from market_data_audit import AuditConfig, audit_market_data, build_run_fingerprints
from optimize import _evaluate_random_window_task, _write_json
from strategy_adapter import resolve_strategy
from runtime_settings import add_runtime_arguments, configure_runtime, runtime_session


def generate_windows(source, start, end, count, min_months, max_months, seed):
    """Sample unique intervals, with actual calendar-month durations."""
    coverage = source.coverage()
    start = pd.Timestamp(start or coverage['first_candle'])
    end = pd.Timestamp(coverage['end_exclusive'] if end == 'latest' else end)
    if start < pd.Timestamp(coverage['first_candle']) or end > pd.Timestamp(coverage['end_exclusive']):
        raise ValueError('Requested range extends outside the dataset')
    if end <= start or start + pd.DateOffset(months=min_months) > end:
        raise ValueError('Range is too short for the minimum window length')
    rng = random.Random(seed)
    times = source.open_times()
    windows, seen = [], set()
    for _ in range(max(10000, count * 200)):
        months = rng.randint(min_months, max_months)
        last_start = end - pd.DateOffset(months=months)
        lower = source.resolve_index(start)
        upper = int(times.searchsorted(last_start, side='right')) - 1
        if upper < lower:
            continue
        first = rng.randint(lower, upper)
        stop = source.resolve_index(times.iloc[first] + pd.DateOffset(months=months))
        if stop <= first or (first, stop) in seen:
            continue
        seen.add((first, stop))
        windows.append(dict(window_id=len(windows) + 1, start=first, end=stop,
                            start_date=source.format_bound(first),
                            end_exclusive=source.format_bound(stop), months=months))
        if len(windows) == count:
            return windows
    raise ValueError(f'Only {len(windows)}/{count} unique windows fit; widen the range or reduce tests')


def summarize(records, expected, min_trades, max_drawdown, min_positive_ratio, min_pass_ratio):
    rows = []
    required = ('total_profit_percent', 'maximum_drawdown', 'closed_trades', 'liquidations')
    for record in records:
        result = record.get('result') or {}
        valid = not record.get('error') and all(
            isinstance(result.get(k), (int, float)) and math.isfinite(result[k]) for k in required)
        passed = bool(valid and result['total_profit_percent'] > 0
                      and abs(result['maximum_drawdown']) <= max_drawdown
                      and result['closed_trades'] >= min_trades and result['liquidations'] == 0)
        rows.append({**{k: v for k, v in record.items() if k != 'result'},
                     **{k: v for k, v in result.items() if not isinstance(v, (list, dict, tuple))},
                     'valid': bool(valid), 'window_pass': passed})
    frame = pd.DataFrame(rows)
    valid = frame[frame['valid']] if len(frame) else frame
    profits = valid['total_profit_percent'] if len(valid) else pd.Series(dtype=float)
    positive = int((profits > 0).sum())
    passed = sum(row['window_pass'] for row in rows)
    gates = dict(complete=len(rows) == expected, no_failed_windows=len(valid) == expected,
                 positive_window_ratio=positive / expected >= min_positive_ratio,
                 passing_window_ratio=passed / expected >= min_pass_ratio)
    summary = dict(status='PASS' if all(gates.values()) else 'FAIL',
                   meaning='Descriptive fixed-parameter robustness check; not independent research acceptance',
                   requested_windows=expected, completed_windows=len(rows),
                   failed_windows=len(rows) - len(valid), positive_window_ratio=positive / expected,
                   passing_window_ratio=passed / expected,
                   median_profit_percent=float(profits.median()) if len(profits) else None,
                   p10_profit_percent=float(profits.quantile(.1)) if len(profits) else None,
                   worst_profit_percent=float(profits.min()) if len(profits) else None,
                   best_profit_percent=float(profits.max()) if len(profits) else None,
                   worst_drawdown_percent=float(valid.maximum_drawdown.abs().max()) if len(valid) else None,
                   gates=gates)
    return summary, frame


def write_workbook(path, frame, summary, params, plan, records):
    """Use the project's existing pandas/openpyxl report stack."""
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    identity = ['window_id', 'start_date', 'end_exclusive', 'months', 'window_pass']
    monthly, trades = [], []
    for record in records:
        for index, value in enumerate((record.get('result') or {}).get('monthly_returns', []), 1):
            monthly.append(dict(window_id=record['window_id'], month_sequence=index, return_fraction=value))
        for index, value in enumerate((record.get('result') or {}).get('trade_profits', []), 1):
            trades.append(dict(window_id=record['window_id'], trade_sequence=index, net_profit=value))
    notes = [
        'Every window starts a fresh account with the same frozen parameters and initial balance.',
        'Windows overlap: do not sum their profits or count repeated trades as independent observations.',
        'PASS/FAIL uses the thresholds on Settings. It is not research accepted=true or a future-profit claim.',
        'Monthly returns are fractions; first and last months can be partial.',
        'Trade Profits contains engine-returned profit sequences, not entry/exit execution logs.',
        'Selection history comes from the adjacent snapshot manifest when available; unknown history is not unseen.',
        'All end boundaries are exclusive. Indicator warmup is enabled and excluded from trading.',
    ]
    scalar_summary = {k: v for k, v in summary.items() if k != 'gates'}
    sheets = {'Summary': pd.DataFrame(scalar_summary.items(), columns=['metric', 'value']),
              'Gates': pd.DataFrame(summary['gates'].items(), columns=['gate', 'passed']),
              'Windows': frame,
              'Parameters': pd.DataFrame(params.items(), columns=['parameter', 'value']),
              'Settings': pd.DataFrame([(k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
                                        for k, v in plan.items() if k not in ('windows', 'params')],
                                       columns=['setting', 'value']),
              'Notes': pd.DataFrame({'note': notes})}
    numeric = frame.loc[frame['valid']].select_dtypes(include='number').drop(
        columns=['window_id', 'start', 'end', 'months', 'duration_s'], errors='ignore')
    if len(numeric) and len(numeric.columns):
        sheets['Metric Statistics'] = numeric.describe(percentiles=[.1, .25, .5, .75, .9]).T.reset_index(
        ).rename(columns={'index': 'metric'})
    for title, prefixes in [('Directional', ('long_', 'short_', 'stronger_side')),
                            ('RSI', ('rsi_',)), ('Scale In', ('scale_',))]:
        sheets[title] = frame[[c for c in frame if c in identity or c.startswith(prefixes)]]
    sheets['Monthly Returns'] = pd.DataFrame(monthly, columns=['window_id', 'month_sequence', 'return_fraction'])
    # Split long trade series rather than silently exceeding the Excel row limit.
    trade_frame = pd.DataFrame(trades, columns=['window_id', 'trade_sequence', 'net_profit'])
    for offset in range(0, max(1, len(trade_frame)), 1_000_000):
        sheets[f'Trade Profits {offset // 1_000_000 + 1}'] = trade_frame.iloc[offset:offset + 1_000_000]
    temporary = path.with_suffix('.tmp.xlsx')
    with pd.ExcelWriter(temporary, engine='openpyxl') as writer:
        for name, data in sheets.items():
            data.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='17365D')
            for column in ws.iter_cols():
                ws.column_dimensions[column[0].column_letter].width = min(65, max(14, max(
                    len(str(cell.value or '')) for cell in column[:100]) + 2))
                for cell in column[1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = '0.00%;[Red]-0.00%' if column[0].value == 'return_fraction' else '#,##0.00;[Red]-#,##0.00'
                    elif isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                        cell.data_type = 's'
            for index, col in enumerate(data.columns, 1):
                if col in ('total_profit_percent', 'net_profit', 'return_fraction') and len(data):
                    letter = get_column_letter(index)
                    ws.conditional_formatting.add(f'{letter}2:{letter}{len(data)+1}', ColorScaleRule(
                        start_type='min', start_color='F8696B', mid_type='percentile', mid_value=50,
                        mid_color='FFEB84', end_type='max', end_color='63BE7B'))
        if 'total_profit_percent' in frame and len(frame):
            buckets = pd.cut(frame.total_profit_percent.dropna(), bins=10).value_counts(sort=False)
            distribution = pd.DataFrame({'return_bucket_percent': [str(x) for x in buckets.index], 'windows': buckets.values})
            distribution.to_excel(writer, sheet_name='Distribution', index=False)
            ws = writer.sheets['Distribution']
            chart = BarChart()
            chart.title = 'Window return distribution (overlapping samples)'
            chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(distribution)+1), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(distribution)+1))
            add_report_chart(writer.book, chart)
    temporary.replace(path)


@runtime_session
def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--params', required=True, help='best_params.json or its containing directory')
    parser.add_argument('--strategy', default='ma')
    parser.add_argument('--tests', type=int, default=1000)
    parser.add_argument('--start', help='Earliest window start; defaults to dataset start')
    parser.add_argument('--end', default='latest', help='Exclusive upper bound')
    parser.add_argument('--min-months', type=int, default=6)
    parser.add_argument('--max-months', type=int, default=12)
    parser.add_argument('--seed', type=int, default=42)
    parser.add_argument('-w', '--workers', type=int, default=8)
    parser.add_argument('--min-trades', type=int, default=5)
    parser.add_argument('--max-drawdown', type=float, default=30)
    parser.add_argument('--min-positive-ratio', type=float, default=.7)
    parser.add_argument('--min-pass-ratio', type=float, default=.7)
    parser.add_argument('--output-dir', default='outputs/evaluate_params')
    parser.add_argument('--resume', action='store_true')
    parser.add_argument('--dry-run', action='store_true')
    add_runtime_arguments(parser)
    args = parser.parse_args(argv)
    try:
        configure_runtime(args, argv)
    except (ValueError, OSError) as exc:
        parser.error(str(exc))
    if min(args.tests, args.workers, args.min_months) <= 0 or args.max_months < args.min_months:
        parser.error('tests/workers/months must be positive and max-months >= min-months')
    if args.min_trades < 0 or not 0 <= args.max_drawdown <= 100 or not all(
            0 <= x <= 1 for x in (args.min_positive_ratio, args.min_pass_ratio)):
        parser.error('Invalid quality thresholds')
    adapter = resolve_strategy(args.strategy)
    source = adapter.market_data_source()
    if source is None:
        parser.error('Strategy must declare DATA_FILE')
    params_path = Path(args.params).resolve()
    if params_path.is_dir():
        params_path /= 'best_params.json'
    params = adapter.load_tune(params_path)
    params = {**adapter.default_values(params), **params}
    if not adapter.validate_candidate(params):
        parser.error('Invalid strategy parameters')
    windows = generate_windows(source, args.start, args.end, args.tests, args.min_months, args.max_months, args.seed)
    provenance_path = params_path.parent / 'manifest.json'
    provenance = json.loads(provenance_path.read_text(encoding='utf-8')) if provenance_path.exists() else {}
    selection_end = provenance.get('development_end_exclusive')
    for window in windows:
        window['selection_history'] = ('overlaps_selection' if pd.Timestamp(window['start_date']) < pd.Timestamp(selection_end)
                                       else 'after_selection') if selection_end else 'unknown'
    plan = {k: v for k, v in vars(args).items() if k not in ('workers', 'performance', 'resume', 'dry_run', 'output_dir')}
    plan.update(params=params, source_params=str(params_path), strategy=adapter.identifier,
                windows=windows, coverage=source.coverage(), selection_end_exclusive=selection_end)
    print(f'Fixed configuration | {len(windows)} windows | {args.min_months}-{args.max_months} calendar months', flush=True)
    print(f"Dataset: {source.coverage()['first_candle']} -> {source.coverage()['end_exclusive']}", flush=True)
    if args.dry_run:
        print(json.dumps(windows[:3], indent=2))
        return
    import openpyxl  # Fail before expensive backtests if the report dependency is absent.
    output = Path(args.output_dir)
    output.mkdir(parents=True, exist_ok=True)
    manifest_path = output / 'manifest.json'
    if not args.resume and any(output.iterdir()):
        parser.error('Output directory is not empty; use a new directory or --resume')
    audit = audit_market_data(source.data_file, AuditConfig(expected_interval=source.interval()))
    audit.raise_for_errors()
    fingerprints = build_run_fingerprints(source.data_file, list(Path(__file__).parent.glob('*.py')),
                                          plan).to_dict()
    manifest = dict(plan=plan, fingerprints=fingerprints)
    if args.resume:
        if not manifest_path.exists() or json.loads(manifest_path.read_text(encoding='utf-8')) != manifest:
            parser.error('Resume mismatch: parameters, data, code or evaluation settings changed')
    else:
        _write_json(manifest_path, manifest)
        _write_json(output / 'market_data_audit.json', audit.to_dict())
        _write_json(output / 'frozen_params.json', params)
    checkpoint = output / 'windows'
    checkpoint.mkdir(exist_ok=True)
    records = [json.loads(p.read_text(encoding='utf-8')) for p in sorted(checkpoint.glob('window_*.json'))]
    done = {r['window_id'] for r in records}
    tasks = [(w['window_id'], 'fixed', w['window_id'], params, w['start'], w['end'], adapter.identifier)
             for w in windows if w['window_id'] not in done]
    by_id = {w['window_id']: w for w in windows}
    started = time.monotonic()
    with multiprocessing.Pool(min(args.workers, max(1, len(tasks)))) as pool:
        for item in pool.imap_unordered(_evaluate_random_window_task, tasks, chunksize=1):
            _, _, window_id, _, _, _, result, elapsed, error = item
            record = dict(by_id[window_id], result=result, duration_s=elapsed, error=error)
            _write_json(checkpoint / f'window_{window_id:06d}.json', record)
            records.append(record)
            if len(records) % 10 == 0 or len(records) == args.tests:
                print(f'{len(records)}/{args.tests} complete | elapsed {time.monotonic()-started:.0f}s', flush=True)
    records.sort(key=lambda r: r['window_id'])
    summary, frame = summarize(records, args.tests, args.min_trades, args.max_drawdown,
                               args.min_positive_ratio, args.min_pass_ratio)
    _write_json(output / 'summary.json', summary)
    frame.to_csv(output / 'window_results.csv', index=False, encoding='utf-8-sig')
    write_workbook(output / 'evaluation_report.xlsx', frame, summary, params, plan, records)
    print(f"{summary['status']} | positive windows {summary['positive_window_ratio']:.1%} | {output / 'evaluation_report.xlsx'}")


if __name__ == '__main__':
    multiprocessing.freeze_support()
    main()
